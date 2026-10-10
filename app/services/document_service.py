
from io import BytesIO
from pathlib import Path

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
import pandas as pd


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".csv",
    ".txt",
    ".pptx",
}


def extract_document_text(
    filename: str,
    file_bytes: bytes
) -> str:

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file format: {extension}"
        )

    # 1. Extract PDF text
    if extension == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))

        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    # 2. Extract Word text
    elif extension == ".docx":
        doc = Document(BytesIO(file_bytes))

        return "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
        )

    # 3. Extract Excel text
    elif extension == ".xlsx":
        workbook = load_workbook(
            BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )

        lines = []

        try:
            for sheet in workbook.worksheets:
                lines.append(f"Sheet: {sheet.title}")

                for row in sheet.iter_rows(values_only=True):
                    values = [
                        str(value)
                        for value in row
                        if value is not None
                    ]

                    if values:
                        lines.append(" | ".join(values))

        finally:
            workbook.close()

        return "\n".join(lines)

    # 4. Extract CSV text
    elif extension == ".csv":
        df = pd.read_csv(BytesIO(file_bytes))

        return df.to_string(index=False)

    # 5. Extract TXT text
    elif extension == ".txt":
        return file_bytes.decode("utf-8-sig")

    # 6. Extract PowerPoint text
    elif extension == ".pptx":
        presentation = Presentation(BytesIO(file_bytes))

        lines = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            lines.append(f"Slide {slide_number}")

            for shape in slide.shapes:

                if shape.has_text_frame:
                    lines.append(shape.text)

                if shape.has_table:
                    for row in shape.table.rows:
                        lines.append(
                            " | ".join(
                                cell.text
                                for cell in row.cells
                            )
                        )

        return "\n".join(lines)

    raise ValueError(
        f"Unsupported file format: {extension}"
    )